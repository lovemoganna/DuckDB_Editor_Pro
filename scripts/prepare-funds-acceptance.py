import argparse,pathlib,hashlib,json
import openpyxl,duckdb
parser=argparse.ArgumentParser(description='Build an independent native DuckDB oracle from only four authorized XLSX sheets.')
parser.add_argument('--source',required=True)
parser.add_argument('--output',default='../audit-20261015')
args=parser.parse_args()
out=pathlib.Path(args.output).resolve();out.mkdir(parents=True,exist_ok=True)
source=pathlib.Path(args.source).resolve()
allowed=['users','fiat_transactions','crypto_transactions','exchange_rates']
book=openpyxl.load_workbook(source,read_only=True,data_only=True)
filtered=openpyxl.Workbook();filtered.remove(filtered.active)
db=duckdb.connect(str(out/'oracle.duckdb'))
data={};schemas={};stats={}
for name in allowed:
    values=iter(book[name].values);headers=list(next(values))
    rows=[list(r[:len(headers)]) for r in values if any(v is not None for v in r)]
    data[name]={'headers':headers,'rows':rows}
    sheet=filtered.create_sheet(name);sheet.append(headers)
    for row in rows:sheet.append(row)
    types=[]
    for i,h in enumerate(headers):
        col=[r[i] for r in rows if r[i] is not None]
        t='DECIMAL(38,12)' if h in ['amount','fee','rate_to_usdt','usdt_to_asset_rate'] else 'BIGINT' if col and all(isinstance(v,(float,int)) and float(v).is_integer() for v in col) else 'VARCHAR'
        types.append(t)
    ddl=','.join('"'+h+'" '+t for h,t in zip(headers,types))
    db.execute(f'CREATE OR REPLACE TABLE "{name}" ({ddl})')
    db.executemany(f'INSERT INTO "{name}" VALUES ({",".join("?" for _ in headers)})',rows)
    schemas[name]=db.execute(f'DESCRIBE "{name}"').fetchall()
    stats[name]={'rows':len(rows),'nulls':{h:sum(r[i] is None for r in rows) for i,h in enumerate(headers)}}
filtered.save(out/'four-tables.xlsx')
queries={
'pivot_count':"SELECT channel, COUNT(*) AS records FROM fiat_transactions GROUP BY channel ORDER BY channel",
'pivot_sum':"SELECT channel,direction,SUM(amount) AS amount FROM fiat_transactions WHERE currency='USD' AND status='completed' GROUP BY channel,direction ORDER BY channel,direction",
'pivot_avg':"SELECT channel,AVG(amount) AS amount FROM fiat_transactions WHERE currency='USD' AND status='completed' GROUP BY channel ORDER BY channel",
'usd_gmv':"SELECT SUM(amount) AS value FROM fiat_transactions WHERE currency='USD' AND status='completed'",
'usd_net':"SELECT SUM(CASE WHEN direction='withdrawal' THEN -amount ELSE amount END) AS value FROM fiat_transactions WHERE currency='USD' AND status='completed'",
'usd_wire':"SELECT SUM(amount) AS value FROM fiat_transactions WHERE currency='USD' AND status='completed' AND channel='wire_transfer'",
'chart':"SELECT channel,SUM(amount) AS amount FROM fiat_transactions WHERE currency='USD' AND status='completed' GROUP BY channel ORDER BY channel",
'user_key':"SELECT COUNT(*) n,COUNT(DISTINCT uid) distinct_keys,COUNT(*) FILTER(WHERE uid IS NULL) null_keys FROM users",
'fiat_key':"SELECT COUNT(*) n,COUNT(DISTINCT fiat_order_id) distinct_keys FROM fiat_transactions",
'crypto_key':"SELECT COUNT(*) n,COUNT(DISTINCT crypto_order_id) distinct_keys FROM crypto_transactions",
'rate_key':"SELECT asset,quote_currency,effective_date,COUNT(*) n FROM exchange_rates GROUP BY ALL HAVING COUNT(*)>1",
'fiat_orphans':"SELECT COUNT(*) n FROM fiat_transactions f LEFT JOIN users u ON CAST(f.uid AS VARCHAR)=u.uid WHERE u.uid IS NULL",
'crypto_orphans':"SELECT COUNT(*) n FROM crypto_transactions f LEFT JOIN users u ON CAST(f.uid AS VARCHAR)=u.uid WHERE u.uid IS NULL",
'fiat_user_cardinality':"SELECT COUNT(*) n FROM fiat_transactions f JOIN users u ON CAST(f.uid AS VARCHAR)=u.uid",
'joined':"SELECT u.country_code,COUNT(*) records,SUM(f.amount) amount FROM fiat_transactions f JOIN users u ON CAST(f.uid AS VARCHAR)=u.uid WHERE f.currency='USD' AND f.status='completed' GROUP BY u.country_code ORDER BY u.country_code",
'rates':"SELECT asset,quote_currency,CAST(effective_date AS VARCHAR) effective_date FROM exchange_rates ORDER BY asset",
'range':"SELECT CAST(MIN(created_at) AS VARCHAR) earliest,CAST(MAX(created_at) AS VARCHAR) latest FROM fiat_transactions",
}
results={}
for key,sql in queries.items():
    cursor=db.execute(sql);results[key]=[dict(zip([c[0] for c in cursor.description],row)) for row in cursor.fetchall()]
report={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'allowed_tables':allowed,'stats':stats,'schemas':schemas,'queries':queries,'results':results,'duckdb_version':duckdb.__version__}
(out/'oracle.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf8')
(out/'source-rows.json').write_text(json.dumps(data,ensure_ascii=False,default=str),encoding='utf8')
db.close()
print(json.dumps({'stats':stats,'results':results},ensure_ascii=False,default=str))
